"""
build_en_pptx.py
================
Export the core English BETEKK CanvasTEKK pitch deck (slides 0-14) from
template.pptx, dropping the appendix (slides 15-41).

The deck ships in English, so this script produces a clean main-pitch copy
and applies safe real edits (contact line / tagline). It targets shapes by
their actual names in the current template.

Usage:
    python output/build_en_pptx.py
"""

import sys
from pathlib import Path

from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")

TEMPLATE = Path(__file__).resolve().parent.parent / "scripts" / "templates" / "template.pptx"
OUTPUT = Path(__file__).resolve().parent / "betekk_pitch_en.pptx"

_REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def delete_slide(prs, index: int) -> None:
    sld_id = prs.slides._sldIdLst[index]
    rid = sld_id.get(f"{{{_REL_NS}}}id")
    if rid:
        prs.part.drop_rel(rid)
    prs.slides._sldIdLst.remove(sld_id)


def find_shape(slide, name: str):
    for shape in slide.shapes:
        if shape.name == name:
            return shape
    return None


def set_first_paragraph(shape, text: str) -> None:
    if not shape or not shape.has_text_frame:
        return
    tf = shape.text_frame
    if tf.paragraphs and tf.paragraphs[0].runs:
        tf.paragraphs[0].runs[0].text = text
        for run in tf.paragraphs[0].runs[1:]:
            run.text = ""
    elif tf.paragraphs:
        tf.paragraphs[0].text = text


def main() -> None:
    if not TEMPLATE.exists():
        print(f"ERROR: template not found: {TEMPLATE}")
        sys.exit(1)

    prs = Presentation(str(TEMPLATE))
    total = len(prs.slides)
    print(f"Loaded template: {total} slides")

    # Keep slides 0-14 (main pitch); delete appendix 15..(total-1) from the top.
    for index in range(total - 1, 14, -1):
        delete_slide(prs, index)
    print(f"Trimmed to {len(prs.slides)} slides (main pitch)")

    slides = list(prs.slides)

    # Safe real edits on the core deck.
    # Cover (S0): tagline.
    tagline = find_shape(slides[0], "tagline")
    set_first_paragraph(tagline, "Every building gets checked thousands of times. We make every check digital.")

    # Closing (S14): contact line.
    sub = find_shape(slides[14], "Subtitle 2")
    set_first_paragraph(sub, "Visit https://betekk.com  |  CanvasTEKK by BETEKK")

    prs.save(str(OUTPUT))
    print(f"Saved: {OUTPUT}")

    # Verify
    prs2 = Presentation(str(OUTPUT))
    print(f"\nVerification - {len(prs2.slides)} slides:")
    for i, slide in enumerate(prs2.slides):
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            txt = "".join(r.text for r in shape.text_frame.paragraphs[0].runs).strip()
            if txt:
                print(f"  Slide {i}: {txt[:70]}")
                break


if __name__ == "__main__":
    main()
